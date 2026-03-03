import os

from openpyxl import Workbook
from openpyxl.styles import PatternFill
import pandas as pd

from custom_prompt_dict import EXTRACT_DEFINITION_DATA


def compare_csv(datetime: str) -> None:
    # CSVファイルの読み込み
    pred_df = pd.read_csv(f"onb_extraction_results/{datetime}_extraction_result.csv")
    annotate_df = pd.read_csv("onb_annotation.csv")

    # 結果を格納するリスト
    results = []

    # ファイル名でマッチング
    for _, ext_row in pred_df.iterrows():
        filename = ext_row["ファイル名"]
        annotate_row = annotate_df[annotate_df["ファイル名"] == filename]

        if len(annotate_row) == 0:
            continue

        annotate_row = annotate_row.iloc[0]
        row_result = {"ファイル名": filename}

        # 各カラムの比較
        headers = ["ファイル名"] + [prompt.japanese_name for prompt in EXTRACT_DEFINITION_DATA.values()]

        for header in headers:
            ext_value = str(ext_row[header])
            annotate_value = str(annotate_row[header])

            # ファイル名はそのまま出力
            if header == "ファイル名":
                row_result[header] = ext_value
            # onb_annotationの値が'-'の場合
            elif annotate_value == "-":
                row_result[header] = ""
            # 値の比較（文字列として扱う）
            else:
                row_result[header] = str(ext_value == annotate_value)

        results.append(row_result)

    # 結果をDataFrameに変換
    result_df = pd.DataFrame(results)

    # 結果を出力するフォルダの作成
    output_dir = f"onb_eval/{datetime}"
    if not os.path.isdir(output_dir):
        os.makedirs(output_dir)

    # 結果をCSVファイルに出力
    result_df.to_csv(f"onb_eval/{datetime}/{datetime}_eval.csv", index=False, encoding="utf-8")


def calculate_accuracy_stats(datetime: str) -> None:
    # 入力CSVの読み込み
    eval_df = pd.read_csv(f"onb_eval/{datetime}/{datetime}_eval.csv")

    # 項目ごとの正答率を計算
    column_stats = []
    for column in eval_df.columns:
        if column == "ファイル名":  # ファイル名列はスキップ
            continue

        # 文字列として扱われているTrue/Falseを確実に検出
        true_count = eval_df[column].astype(str).str.strip().eq("True").sum()
        false_count = eval_df[column].astype(str).str.strip().eq("False").sum()
        total_count = true_count + false_count

        if total_count > 0:
            accuracy = true_count / total_count
            column_stats.append(
                {
                    "項目名": column,
                    "正答率": round(accuracy, 2),
                    "正解数": true_count,
                    "項目数": total_count,
                }
            )

    # ファイルごとの正答率を計算
    file_stats = []
    all_true_count = 0
    all_total_count = 0

    for _, row in eval_df.iterrows():
        # ファイル名列を除外して、True/Falseの値のみを対象とする
        values = row[1:].astype(str).str.strip()  # 文字列に変換してスペースを削除
        true_count = (values == "True").sum()
        false_count = (values == "False").sum()
        total_count = true_count + false_count

        all_true_count += true_count
        all_total_count += total_count

        if total_count > 0:
            accuracy = true_count / total_count
            file_stats.append(
                {
                    "ファイル名": row["ファイル名"],
                    "正答率": round(accuracy, 2),
                    "正解数": true_count,
                    "項目数": total_count,
                }
            )

    # 全体の精度を計算
    overall_accuracy = all_true_count / all_total_count if all_total_count > 0 else 0
    overall_stats = [
        {
            "全体の精度": round(overall_accuracy, 2),
            "正解項目数": all_true_count,
            "抽出項目数": all_total_count,
        }
    ]

    # 結果をDataFrameに変換し、空の場合でもカラムを持つDataFrameを作成
    column_stats_df = (
        pd.DataFrame(column_stats) if column_stats else pd.DataFrame(columns=["項目名", "正答率", "正解数", "項目数"])
    )
    file_stats_df = (
        pd.DataFrame(file_stats) if file_stats else pd.DataFrame(columns=["ファイル名", "正答率", "正解数", "項目数"])
    )
    overall_stats_df = pd.DataFrame(overall_stats)

    # CSVファイルとして出力
    column_stats_df.to_csv(
        f"onb_eval/{datetime}/{datetime}_eval_column_stats.csv",
        index=False,
        encoding="utf-8",
    )
    file_stats_df.to_csv(
        f"onb_eval/{datetime}/{datetime}_eval_file_stats.csv",
        index=False,
        encoding="utf-8",
    )
    overall_stats_df.to_csv(f"onb_eval/{datetime}/{datetime}_eval_all.csv", index=False, encoding="utf-8")


def read_data(datetime: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    pred_df = pd.read_csv(f"onb_extraction_results/{datetime}_extraction_result.csv")
    annotate_df = pd.read_csv("onb_annotation.csv")
    return pred_df, annotate_df


def compare_data(
    pred_df: pd.DataFrame, annotate_df: pd.DataFrame
) -> tuple[list[dict[str, str]], list[tuple[int, int]]]:
    results = []
    cells_to_highlight = []
    headers = ["ファイル名"] + [prompt.japanese_name for prompt in EXTRACT_DEFINITION_DATA.values()]

    for row_idx, ext_row in enumerate(pred_df.iterrows()):
        filename = ext_row[1]["ファイル名"]
        annotate_row = annotate_df[annotate_df["ファイル名"] == filename]

        if len(annotate_row) == 0:
            continue

        annotate_row = annotate_row.iloc[0]
        row_result = {"ファイル名": filename}

        for col_idx, header in enumerate(headers):
            ext_value = str(ext_row[1][header])
            annotate_value = str(annotate_row[header])

            if header == "ファイル名":
                row_result[header] = ext_value
            elif annotate_value == "-":
                row_result[header] = ""
            elif ext_value == annotate_value:
                row_result[header] = f"True: {annotate_value}"
            else:
                row_result[header] = f"False: 正解:{annotate_value} / 抽出:{ext_value}"
                cells_to_highlight.append((row_idx + 2, col_idx + 1))

        results.append(row_result)

    return results, cells_to_highlight


def create_excel_file(results: list[dict[str, str]], cells_to_highlight: list[tuple[int, int]], datetime: str) -> str:
    result_df = pd.DataFrame(results)
    output_dir = f"onb_eval/{datetime}"
    if not os.path.isdir(output_dir):
        os.makedirs(output_dir)

    wb = Workbook()
    ws = wb.active

    for col_idx, header in enumerate(result_df.columns, 1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, row in enumerate(result_df.values, 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    for row, col in cells_to_highlight:
        ws.cell(row=row, column=col).fill = yellow_fill

    for column in ws.columns:
        max_length = max(len(str(cell.value)) for cell in column)
        adjusted_width = max_length + 2
        ws.column_dimensions[column[0].column_letter].width = adjusted_width

    excel_path = f"{output_dir}/{datetime}_eval.xlsx"
    wb.save(excel_path)
    return excel_path


def compare_and_export_excel(datetime: str) -> str:
    pred_df, annotate_df = read_data(datetime)
    results, cells_to_highlight = compare_data(pred_df, annotate_df)
    excel_path = create_excel_file(results, cells_to_highlight, datetime)
    return excel_path


if __name__ == "__main__":
    # onb_extraction_resultsのdatetimeを設定!!
    datetime = "2026-03-03-07-27"
    compare_csv(datetime)
    calculate_accuracy_stats(datetime)
    compare_and_export_excel(datetime)
